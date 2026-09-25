# build_docx.py - proposal md -> .docx (headings, paragraphs, real tables).
# Usage: py proposal/build_docx.py proposal/HOBYD-Proposal-Updated.md proposal/HOBYD-Proposal-Updated.docx
import re
import sys
from docx import Document
from docx.shared import Pt

SRC, DST = sys.argv[1], sys.argv[2]
lines = open(SRC, encoding="utf-8").read().splitlines()

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

i, tables, paras = 0, 0, 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1
        continue
    if ln.startswith("### "):
        doc.add_heading(ln[4:], level=3)
        paras += 1
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=2)
        paras += 1
    elif ln.startswith("# "):
        t = doc.add_heading(ln[2:], level=1)
        paras += 1
    elif ln.strip().startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            # skip separator rows (|---|---|)
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                i += 1
                continue
            rows.append(cells)
            i += 1
        if rows:
            width = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=width)
            tbl.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci in range(width):
                    tbl.cell(ri, ci).text = r[ci] if ci < len(r) else ""
            tables += 1
        continue
    elif ln.startswith("[Gambar") or ln.startswith("["):
        p = doc.add_paragraph()
        r = p.add_run(ln)
        r.italic = True
        paras += 1
    else:
        doc.add_paragraph(ln)
        paras += 1
    i += 1

doc.save(DST)
print(f"saved {DST}: {paras} paragraphs/headings, {tables} tables")
