# update_proposal.py - bedah teks docx asli (clone-and-edit).
# Tidak menyentuh gambar, caption (kecuali Gbr 4), tabel selain B/D, style, bullet.
# Usage: py proposal/update_proposal.py "<path docx asli>" "<path output>"
import re
import sys
from copy import deepcopy
from docx import Document

SRC, DST = sys.argv[1], sys.argv[2]
doc = Document(SRC)

def find(prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    return hits

def replace_once(prefix, new):
    hits = find(prefix)
    assert len(hits) == 1, f"anchor x{len(hits)}: {prefix[:50]}"
    p = hits[0]
    assert len(p.runs) <= 1, f"multi-run: {prefix[:50]}"
    p.text = new
    print(f"OK replace: {prefix[:45]}")

def insert_after(prefix, texts):
    hits = find(prefix)
    assert len(hits) == 1, f"anchor x{len(hits)}: {prefix[:50]}"
    anchor = hits[0]
    ref = anchor._p
    for t in texts:
        el = deepcopy(anchor._p)
        el.clear_content() if hasattr(el, "clear_content") else None
        # fresh paragraph with same style
        from docx.oxml import OxmlElement
        new_p = OxmlElement("w:p")
        ref.addnext(new_p)
        from docx.text.paragraph import Paragraph
        np = Paragraph(new_p, anchor._parent)
        np.style = anchor.style
        np.add_run(t)
        ref = new_p
    print(f"OK insert x{len(texts)} after: {prefix[:45]}")

# ---- TOC: buang nomor halaman ----
n_toc = 0
for p in doc.paragraphs:
    if re.match(r"^\d+\.\s", p.text) and "." * 5 in p.text and p.style.name == "Normal":
        p.text = re.sub(r"\s*\.+\s*\d+\s*$", "", p.text)
        n_toc += 1
print(f"OK toc stripped: {n_toc}")

# ---- §1 ----
replace_once(
    "Hobyd (Hobby and Bid) adalah platform",
    "Hobyd (Hobby and Bid) adalah platform live auction berbasis website untuk kolektor. "
    "Penjual menyiarkan live dan melelang item satu per satu dengan timer. Pembeli mengajukan bid "
    "secara real time dan melihat harga terbentuk melalui kompetisi yang transparan. Pembedanya: "
    "pengalaman yang dirancang untuk kolektor, penemuan harga real time, mesin lelang yang otoritatif "
    "di server, riwayat bid yang terbuka, dan lapisan kepercayaan transaksi.",
)
replace_once(
    "Target awal adalah kolektor Pokemon TCG",
    "Target awal adalah kolektor Pokemon TCG di Indonesia. Pada hackathon, tim mendemokan satu room live "
    "dengan satu penjual dan beberapa penawar: bid real time dengan validasi server, perpanjangan anti-sniping, "
    "penguncian pemenang, dan demo pembayaran. Semua alur ini berjalan pada MVP yang dibangun: penjual membuat "
    "live dari dialog bernama, memilih kategori (termasuk Other), melelang item dengan mode soft close atau sudden "
    "death, dan pemenang membayar lalu bertukar kontak WhatsApp untuk serah terima. Kategori lain menyusul setelah "
    "model inti tervalidasi.",
)

# ---- §2 ----
replace_once(
    "Momen ini tepat karena tiga kebiasaan",
    "Momen ini tepat karena tiga kebiasaan sudah terbentuk: menonton konten live, membayar dengan dompet digital, "
    "dan berkumpul di komunitas digital. Pengujian live dengan dua ponsel menemukan dua friksi yang tidak terlihat "
    "di desktop: kamera ponsel yang publish di atas 720p tidak selalu bisa di-decode ponsel penawar, dan dialog "
    "yang tidak di-portal bisa terpotong navigasi bawah. Keduanya diperbaiki di MVP. Yang belum tersedia dalam satu "
    "pengalaman yang fokus pada lelang adalah penemuan barang secara live, bidding terstruktur, dan alur transaksi kolektor.",
)

# ---- §4 ----
replace_once(
    "Ruang lelang live Hobyd menampilkan",
    "Ruang lelang live Hobyd menampilkan video penjual, kartu item, harga berjalan, countdown, feed bid, dan tombol "
    "BID dalam satu layar. Loop pembeli: DISCOVER, WATCH, BID, WIN, PAY. Loop penjual: CREATE LIVE, ADD ITEM, START "
    "AUCTION, ACCEPT BIDS, CLOSE AUCTION.",
)
insert_after(
    "Ruang lelang live Hobyd menampilkan",
    [
        "Video live: siaran penjual yang sedang melelang. Penjual dari ponsel mendapat pratinjau kamera portrait 9:16; "
        "penonton ponsel mendapat video full-bleed dengan kartu bid dan chat sebagai overlay; penonton desktop mendapat "
        "side rail berisi kartu bid di atas dan chat di bawah, plus mode theater dan fullscreen dengan rail yang sama.",
        "Kartu item memakai foto 16:9 yang di-crop otomatis saat diunggah, sehingga pratinjau sebelum posting sama "
        "persis dengan yang dilihat pembeli di lobi, kartu item, rail, mobile, dan sales.",
        "Pembuatan live: dialog Go Live meminta judul (1–80 karakter) dan kategori di awal, sehingga tidak ada room "
        "tanpa nama. Kategori: Sneakers, TCG, Vintage Clothing, Electronics, Other. Pemilik room dapat mengganti judul "
        "dan kategori belakangan.",
    ],
)

# ---- §5 ----
replace_once(
    "Inovasi produk:",
    "Inovasi produk: pengalaman yang dirancang untuk kolektor, penemuan barang secara live, penemuan harga real time, "
    "lelang per item yang terstruktur dengan dua mode (soft close yang memperpanjang timer dan sudden death yang tidak "
    "pernah bergeser), lapisan kepercayaan, dan alur transaksi Indonesia. Inovasi engineering: state lelang yang "
    "otoritatif di server. Browser peserta tidak menentukan hasil lelang.",
)

# ---- §6 ----
replace_once(
    "Penjual memulai live dan menampilkan item pertama",
    "Penjual membuka dialog Go Live, memberi judul dan kategori, lalu menyalakan kamera dan memulai stream.",
)
replace_once(
    "Penawar bergabung, menonton siaran, dan menekan BID.",
    "Penjual menampilkan item pertama beserta harga awal, lalu menekan Start bid untuk menjalankan timer. Penawar "
    "bergabung, menonton siaran, dan menekan BID.",
)
replace_once(
    "Timer habis, pemenang dikunci dan diumumkan, lalu masuk demo pembayaran.",
    "Timer habis atau penjual menutup lebih awal (konfirmasi dua ketuk); pemenang dikunci dan diumumkan, lalu masuk "
    "demo pembayaran.",
)
insert_after(
    "Timer habis atau penjual menutup lebih awal",
    [
        "Pemenang membayar dalam jendela 5:00, lalu meninggalkan nomor WhatsApp di halaman menang; penjual melihat "
        "nomor itu di halaman yang sama dan di inbox Sales, lalu menghubungi untuk serah terima.",
    ],
)
replace_once(
    "Alur pembayaran menangani dua hasil:",
    "Alur pembayaran menangani dua hasil: lunas dalam jendela waktu menjadi order confirmed; kedaluwarsa menjadi "
    "expired dan item di-relist (kembali dilelang) sebagai item baru. Kriteria sukses demo: banyak penawar dapat bid "
    "bersamaan, semua klien menerima state yang konsisten, bid tidak valid ditolak server, bid di akhir memperpanjang "
    "timer, pemenang tidak berubah setelah lock, dan pemenang dapat masuk alur pembayaran. Kriteria ini diuji pada "
    "pengujian live dua ponsel plus desktop.",
)

# ---- §10 ----
replace_once(
    "Otorisasi: hanya peserta room yang dapat bid; hanya penjual yang mengontrol lelang.",
    "Otorisasi: hanya peserta room yang dapat bid; hanya pemilik room yang mengontrol lelang; hanya pemenang yang "
    "dapat menyimpan kontak serah terima; daftar order penjual dibatasi room miliknya sendiri.",
)
insert_after(
    "Isolasi kegagalan antara lapisan video dan state lelang.",
    [
        "Satu topik realtime per komponen; tidak ada dua subscriber pada satu topik.",
        "Pembersih otomatis: room preview yang tidak pernah live lebih dari 10 menit terhapus saat lobi dibuka; "
        "room live yang tidak berpenghuni melewati batas yang sama diakhiri (ended), bukan dihapus, agar order tetap tersimpan.",
    ],
)

# ---- §11 ----
replace_once(
    "MVP hackathon: satu penjual, beberapa penonton,",
    "MVP hackathon yang dibangun: live satu penjual dan beberapa penonton; dialog Go Live bernama dengan kategori; "
    "video live (portrait penjual mobile, full-bleed penonton mobile, side rail + theater + fullscreen desktop); kartu "
    "item dengan foto 16:9 WYSIWYG; bidding real time dengan validasi server; timer dengan mode soft close dan sudden "
    "death; kategori gate termasuk Other; anti-sniping; winner lock; demo pembayaran 5:00; kontak WhatsApp pemenang; "
    "inbox Sales; pembersih room terbengkalai. Suite 90 tes unit lolos, typecheck bersih, production build hijau.",
)
replace_once(
    "Di luar MVP: verifikasi penjual lanjutan,",
    "Di luar MVP: verifikasi penjual lanjutan, reputasi, pengungkapan kondisi, penanganan sengketa, logistik, analitik, "
    "chat order dalam aplikasi, dan kategori tambahan. Tahapan: Fase 1 pilot komunitas Pokemon TCG (KPI: retensi penjual, "
    "partisipasi ulang) — berjalan; Fase 2 validasi product-market fit (KPI: penyelesaian lelang dan pembayaran); Fase 3 "
    "perbaikan trust dan transaksi; Fase 4 ekspansi kategori. Metrik produk: auction completion, invalid bid rate, payment "
    "completion, sinkronisasi dan latensi. Metrik marketplace: retensi penjual, partisipasi ulang, item terjual per live, "
    "lelang selesai per live.",
)

# ---- §12 ----
replace_once(
    "Transaksi kolektor membutuhkan penemuan harga",
    "Transaksi kolektor membutuhkan penemuan harga yang transparan dalam alur yang terstruktur. Hobyd menjawabnya dengan "
    "live auction untuk kolektor yang didukung mesin lelang server-authoritative dan riwayat bid terbuka. Demo hackathon "
    "menunjukkan live, bid, win, dan pay berjalan, ditambah serah terima kontak dan inbox Sales. Langkah berikut: pilot "
    "pada komunitas Pokemon TCG untuk memvalidasi retensi dan partisipasi sebelum memperbaiki dan berekspansi.",
)

# ---- caption Gbr 4 ----
replace_once(
    "Gambar 4. Mockup ruang lelang Hobyd",
    "Gambar 4. Ruang lelang Hobyd (desktop dan mobile).",
)

# ---- tabel: model data + risiko ----
for t in doc.tables:
    heads = [c.text for c in t.rows[0].cells]
    if heads[:2] == ["Tabel", "Kolom kunci"]:
        for row in t.rows:
            if row.cells[0].text == "rooms":
                row.cells[1].text = ("id, title, seller_name, owner_id, status (preview/live/ended), category "
                                     "(Sneakers/TCG/Vintage Clothing/Electronics/Other), thumbnail_url")
            elif row.cells[0].text == "items":
                row.cells[1].text = ("id, room_id, title, img_url, start_price, current_price, ends_at, status, "
                                     "auction_mode (soft/hard), duration_sec (10–300), winner")
            elif row.cells[0].text == "orders":
                row.cells[1].text = ("id, item_id, winner, winner_contact, status (pending/paid/expired/cancelled), paid_at")
        # tambah baris chat_messages
        row = t.add_row()
        row.cells[0].text = "chat_messages"
        row.cells[1].text = "id, room_id, nickname, body, created_at"
        print("OK tabel model data")
    elif heads[:3] == ["Risiko", "Dampak", "Mitigasi"]:
        for vals in [
            ("Ponsel tidak bisa decode stream ponsel", "Layar hitam, chat tetap jalan",
             "Capture dibatasi 720p; flag ?debug=1 untuk repro"),
            ("Dialog terpotong navigasi bawah", "Penjual mobile tidak bisa Go Live",
             "Dialog di-portal ke document body"),
        ]:
            row = t.add_row()
            for c, v in zip(row.cells, vals):
                c.text = v
        print("OK tabel risiko +2")

doc.save(DST)
print("saved", DST)
