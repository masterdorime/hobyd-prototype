# swap_images.py - ganti 4 figur + tambah Gambar 18 pada proposal final.
# Usage: py proposal/swap_images.py
import io
from PIL import Image
from docx import Document
from docx.shared import Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC = "proposal/HOBYD-Proposal-Final-Updated.docx"
IMG = "proposal/img"

def load(n):
    im = Image.open(f"{IMG}/{n}")
    return im.convert("RGB")

def side_by_side(a, b, h=880, gap=28):
    def fit(im):
        w = round(im.width * h / im.height)
        return im.resize((w, h))
    A, B = fit(a), fit(b)
    canvas = Image.new("RGB", (A.width + gap + B.width, h), "white")
    canvas.paste(A, (0, 0))
    canvas.paste(B, (A.width + gap, 0))
    return canvas

def to_jpg(im):
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    return buf.getvalue()

def to_png(im):
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()

def blip_info(doc, media_name):
    """(paragraph, extent_el, part) untuk media file tertentu."""
    for p in doc.paragraphs:
        for blip in p._p.xpath(".//a:blip"):
            rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = doc.part.related_parts[rid]
            if part.partname.endswith("/" + media_name):
                extent = p._p.xpath(".//wp:extent")[0]
                return p, extent, part
    raise LookupError(media_name)

def swap(doc, media_name, data):
    p, extent, part = blip_info(doc, media_name)
    old_cx, old_cy = int(extent.get("cx")), int(extent.get("cy"))
    im = Image.open(io.BytesIO(data))
    a = im.height / im.width
    if old_cx * a <= old_cy:
        cx, cy = old_cx, round(old_cx * a)
    else:
        cy, cx = old_cy, round(old_cy / a)
    part._blob = data
    extent.set("cx", str(cx))
    extent.set("cy", str(cy))
    print(f"OK {media_name}: {len(data)//1024}KB frame {cx}x{cy}emu")
    return cx

doc = Document(DOC)

g4 = side_by_side(load("tampilan penonton desktop.png"), load("tampilan live mobile.jpeg"))
swap(doc, "image6.png", to_png(g4))

g15 = side_by_side(load("tampilan home page.png"), load("tampilan home mobile.jpeg"))
swap(doc, "image16.jpeg", to_jpg(g15))

swap(doc, "image17.jpeg", to_jpg(load("tampilan pemenang.png")))
swap(doc, "image18.jpeg", to_jpg(load("tampilan pembayaran sandbox.png")))

# caption Gbr 15 menyesuaikan komposit
for p in doc.paragraphs:
    if p.text.startswith("Gambar 15."):
        p.text = "Gambar 15. Halaman utama (desktop dan mobile)."
        print("OK caption 15")

# Gambar 18 baru: tiru lebar bingkai Gbr 17
_, extent17, _ = blip_info(doc, "image18.jpeg")
w17 = Emu(int(extent17.get("cx")))
caps = [p for p in doc.paragraphs if p.text.startswith("Gambar 17.")]
assert len(caps) == 1
anchor = caps[0]._p
streamer = load("tampilan live streamer.png")
buf = io.BytesIO()
streamer.save(buf, "JPEG", quality=88)
buf.seek(0)

from docx.oxml import OxmlElement
new_p = OxmlElement("w:p")
anchor.addnext(new_p)
from docx.text.paragraph import Paragraph
ip = Paragraph(new_p, caps[0]._parent)
ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
ip.add_run().add_picture(buf, width=w17)

cap_p = OxmlElement("w:p")
new_p.addnext(cap_p)
cp = Paragraph(cap_p, caps[0]._parent)
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
cp.add_run("Gambar 18. Tampilan penjual live.")
print("OK gambar 18")

doc.save(DOC)
print("saved", DOC)
